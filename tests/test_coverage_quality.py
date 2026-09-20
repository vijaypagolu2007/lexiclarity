"""Focused regression tests for production modules covered by the CI gate."""

import io
import json
from types import SimpleNamespace
from unittest.mock import MagicMock

import fitz
import pytest
from docx import Document

from core.parser import process_uploaded_file
from core.parser import truncate as parser_truncate
from core.retrieval import get_relevant_chunks
from src import llm
from src.extract import ExtractionError, extract_text
from src.grounding import check_items, grounded_rate, span_in_document
from src.llm import LLMError, _parse_json, load_prompt, run_task
from src.pdf_export import build_lawyer_prep_pdf, build_simplified_pdf
from src.retrieval import cosine_similarity


def test_parser_handles_text_docx_and_limits():
    upload = SimpleNamespace(name="contract.txt", type="text/plain")
    text, clauses = process_uploaded_file(upload, b"Rent is due monthly.\n\nThis clause is long enough to be retained.")
    assert "Rent is due" in text
    assert len(clauses) == 1
    assert parser_truncate("abcdef", 3).startswith("abc")

    document = Document()
    document.add_paragraph("A sufficiently long DOCX paragraph for extraction.")
    buffer = io.BytesIO()
    document.save(buffer)
    extracted, _ = process_uploaded_file(SimpleNamespace(name="contract.docx", type=""), buffer.getvalue())
    assert "sufficiently long" in extracted

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A sufficiently long PDF clause for extraction.")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    parsed_pdf, _ = process_uploaded_file(SimpleNamespace(name="contract.pdf", type="application/pdf"), pdf_bytes)
    assert "sufficiently long PDF" in parsed_pdf


def test_extract_text_supports_text_and_docx_and_rejects_invalid():
    assert extract_text("notice.txt", b"This is a sufficiently long plain text notice.").startswith("This")
    document = Document()
    document.add_paragraph("A sufficiently long paragraph in a document.")
    buffer = io.BytesIO()
    document.save(buffer)
    assert "paragraph" in extract_text("notice.docx", buffer.getvalue())
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_text("notice.csv", b"data")
    with pytest.raises(ExtractionError, match="empty"):
        extract_text("notice.txt", b"")


def test_llm_parsing_and_successful_run(monkeypatch):
    assert _parse_json("```json\n{\"ok\": true}\n```") == {"ok": True}
    with pytest.raises(LLMError, match="top-level"):
        _parse_json("[]")
    with pytest.raises(LLMError, match="Prompt template"):
        load_prompt("missing-template")

    response = SimpleNamespace(text=json.dumps({"answer": "yes"}), usage_metadata=None)
    client = SimpleNamespace(models=SimpleNamespace(generate_content=MagicMock(return_value=response)))
    assert run_task("chat", "question: test", client=client) == {"answer": "yes"}
    llm._last_metrics = {}


def test_llm_errors_are_actionable(monkeypatch):
    client = SimpleNamespace(models=SimpleNamespace(generate_content=MagicMock(side_effect=ConnectionError())))
    monkeypatch.setattr("src.llm.time.sleep", lambda _: None)
    with pytest.raises(LLMError, match="network connection"):
        run_task("chat", "question: test", client=client)


def test_llm_rejects_empty_and_malformed_responses():
    with pytest.raises(LLMError, match="empty response"):
        _parse_json("")
    with pytest.raises(LLMError, match="invalid JSON"):
        _parse_json("{broken")


def test_grounding_and_semantic_helpers():
    assert span_in_document("rent is due", "Rent is due monthly")
    items = check_items([{"source_span": "Rent is due"}, {"source_span": "invented"}], "Rent is due monthly")
    assert grounded_rate(items) == 0.5
    assert cosine_similarity([1, 0], [1, 0]) == 1.0
    assert cosine_similarity([0, 0], [1, 0]) == 0.0


def test_retrieval_falls_back_when_embeddings_are_unavailable(monkeypatch):
    monkeypatch.setattr("core.retrieval.get_embedding", MagicMock(side_effect=ConnectionError()))
    result = get_relevant_chunks("rent", ["Rent is due monthly.", "The premises are residential."], top_k=1)
    assert result == "Rent is due monthly."


def test_retrieval_ranks_semantic_matches(monkeypatch):
    vectors = {
        "rent": [1.0, 0.0],
        "Rent is due monthly.": [1.0, 0.0],
        "The premises are residential.": [0.0, 1.0],
    }
    monkeypatch.setattr("core.retrieval.get_embedding", lambda text: vectors[text])
    assert get_relevant_chunks("rent", list(vectors)[1:], top_k=1) == "Rent is due monthly."


def test_llm_empty_model_response_is_actionable():
    client = SimpleNamespace(models=SimpleNamespace(generate_content=MagicMock(return_value=SimpleNamespace(text=""))))
    with pytest.raises(LLMError, match="empty response"):
        run_task("chat", "question: test", client=client)


def test_llm_missing_api_key_is_actionable(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    llm._client = None
    with pytest.raises(LLMError, match="No API key"):
        llm._get_client()


def test_pdf_exports_are_real_documents():
    simplified = build_simplified_pdf("Lease", "Simple", [{"original_heading": "Rent", "plain_text": "Rent is due.", "source_span": "Rent is due."}])
    lawyer = build_lawyer_prep_pdf({"case_summary": "Summary", "parties": [], "important_dates": [], "financial_obligations": [], "top_risks": [], "missing_or_ambiguous": [], "questions_for_lawyer": [], "documents_to_bring": [], "timeline": []})
    assert simplified.startswith(b"%PDF")
    assert lawyer.startswith(b"%PDF")
